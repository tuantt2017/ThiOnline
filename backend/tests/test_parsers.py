import os
import tempfile
import pytest
import pymupdf
import docx

from app.models.document import DocumentType, KnowledgeNodeType
from app.services.parser import (
    PDFParser,
    DOCXParser,
    TXTParser,
    parse_document_file,
    DocumentParsingError,
)
from app.services.extractor import StructureExtractor


SAMPLE_TEXTBOOK_CONTENT = """
CHƯƠNG 1: MỆNH ĐỀ VÀ TẬP HỢP

BÀI 1: MỆNH ĐỀ TOÁN HỌC

I. Mệnh đề và mệnh đề chứa biến
1. Khái niệm mệnh đề
Khái niệm: Mệnh đề toán học là một khẳng định đúng hoặc một khẳng định sai.
Một khẳng định không thể vừa đúng vừa sai.

Mục tiêu: Học sinh nhận biết được mệnh đề toán học, phân biệt được mệnh đề đúng và mệnh đề sai.

2. Mệnh đề chứa biến
Định nghĩa: Cho câu P(n) chứa biến n thuộc tập X. Với mỗi giá trị cụ thể của n trong X, P(n) trở thành một mệnh đề.
Ghi nhớ: Mệnh đề chứa biến không phải là mệnh đề nếu chưa gán giá trị cụ thể cho biến.

II. Mệnh đề phủ định
Khái niệm: Phủ định của mệnh đề P là mệnh đề không phải P, kí hiệu là P ngang.
Yêu cầu cần đạt: Học sinh biết cách lập mệnh đề phủ định của một mệnh đề đã cho.
"""


def test_txt_parser_valid():
    with tempfile.NamedTemporaryFile("w", encoding="utf-8", suffix=".txt", delete=False) as f:
        f.write(SAMPLE_TEXTBOOK_CONTENT)
        temp_path = f.name

    try:
        parsed = TXTParser.parse(temp_path)
        assert parsed.file_type == DocumentType.TXT
        assert parsed.total_pages >= 1
        assert parsed.total_chars > 200
        assert "MỆNH ĐỀ VÀ TẬP HỢP" in parsed.raw_text
    finally:
        if os.path.exists(temp_path):
            os.remove(temp_path)


def test_docx_parser_valid():
    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as f:
        temp_path = f.name

    try:
        doc = docx.Document()
        doc.add_heading("CHƯƠNG 1: MỆNH ĐỀ VÀ TẬP HỢP", level=1)
        doc.add_heading("BÀI 1: MỆNH ĐỀ TOÁN HỌC", level=2)
        doc.add_paragraph("Khái niệm: Mệnh đề toán học là một khẳng định đúng hoặc sai.")
        doc.add_paragraph("Mục tiêu: Học sinh nhận biết được mệnh đề toán học.")
        doc.save(temp_path)

        parsed = DOCXParser.parse(temp_path)
        assert parsed.file_type == DocumentType.DOCX
        assert parsed.total_pages >= 1
        assert "MỆNH ĐỀ VÀ TẬP HỢP" in parsed.raw_text
        assert "Khái niệm:" in parsed.raw_text
    finally:
        if os.path.exists(temp_path):
            os.remove(temp_path)


def test_pdf_parser_valid():
    with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as f:
        temp_path = f.name

    try:
        doc = pymupdf.open()
        page1 = doc.new_page()
        page1.insert_text((50, 72), "CHƯƠNG 1: Menh de va tap hop\nBAI 1: Menh de toan hoc\nKhai niem: Menh de la khang dinh dung hoac sai.", fontname="helv")
        
        page2 = doc.new_page()
        page2.insert_text((50, 72), "Muc tieu: Hoc sinh hieu dinh nghia menh de toan hoc.\nGhi nho: Khong the vua dung vua sai.", fontname="helv")
        doc.save(temp_path)
        doc.close()

        parsed = PDFParser.parse(temp_path)
        assert parsed.file_type == DocumentType.PDF
        assert parsed.total_pages == 2
        assert "Menh de va tap hop" in parsed.pages[0].text
        assert "Muc tieu" in parsed.pages[1].text
    finally:
        if os.path.exists(temp_path):
            os.remove(temp_path)



def test_parser_empty_or_corrupt_files():
    # Non-existent file
    with pytest.raises(DocumentParsingError):
        parse_document_file("non_existent_file.pdf", DocumentType.PDF)

    # Empty TXT file
    with tempfile.NamedTemporaryFile("w", encoding="utf-8", suffix=".txt", delete=False) as f:
        f.write("")
        empty_txt = f.name

    try:
        with pytest.raises(DocumentParsingError):
            TXTParser.parse(empty_txt)
    finally:
        if os.path.exists(empty_txt):
            os.remove(empty_txt)


def test_structure_extractor_and_8_level_knowledge_map():
    with tempfile.NamedTemporaryFile("w", encoding="utf-8", suffix=".txt", delete=False) as f:
        f.write(SAMPLE_TEXTBOOK_CONTENT)
        temp_path = f.name

    try:
        parsed = TXTParser.parse(temp_path)
        extractor = StructureExtractor(
            subject="Toán học",
            grade=10,
            book_series="Kết nối tri thức",
            title="Sách Giáo Khoa Toán 10 Tập 1",
        )
        result = extractor.extract(parsed)

        # Verify Chunks
        assert len(result.chunks) > 0
        first_chunk = result.chunks[0]
        assert first_chunk.chapter is not None
        assert "Chương 1" in first_chunk.chapter
        assert first_chunk.lesson is not None
        assert "Bài 1" in first_chunk.lesson
        assert first_chunk.char_count > 0

        # Verify 8-level Knowledge Map hierarchy
        assert len(result.root_nodes) == 1
        subject_node = result.root_nodes[0]
        assert subject_node.node_type == KnowledgeNodeType.SUBJECT
        assert subject_node.title == "Toán học"

        # Level 2: Grade
        assert len(subject_node.children) == 1
        grade_node = subject_node.children[0]
        assert grade_node.node_type == KnowledgeNodeType.GRADE
        assert "10" in grade_node.title

        # Level 3: Book
        assert len(grade_node.children) == 1
        book_node = grade_node.children[0]
        assert book_node.node_type == KnowledgeNodeType.BOOK
        assert "Kết nối tri thức" in book_node.title

        # Level 4: Chapter
        assert len(book_node.children) >= 1
        chapter_node = book_node.children[0]
        assert chapter_node.node_type == KnowledgeNodeType.CHAPTER
        assert "Chương 1" in chapter_node.title

        # Level 5: Lesson
        assert len(chapter_node.children) >= 1
        lesson_node = chapter_node.children[0]
        assert lesson_node.node_type == KnowledgeNodeType.LESSON
        assert "Bài 1" in lesson_node.title

        # Level 6: Topic
        assert len(lesson_node.children) >= 1
        topic_node = lesson_node.children[0]
        assert topic_node.node_type == KnowledgeNodeType.TOPIC

        # Level 7: Concept
        assert len(topic_node.children) >= 1
        concept_node = topic_node.children[0]
        assert concept_node.node_type == KnowledgeNodeType.CONCEPT

        # Level 8: Learning Objective
        assert len(concept_node.children) >= 1
        objective_node = concept_node.children[0]
        assert objective_node.node_type == KnowledgeNodeType.LEARNING_OBJECTIVE

        # Verify summary metadata
        assert result.summary_metadata["total_chapters"] >= 1
        assert result.summary_metadata["total_lessons"] >= 1
        assert result.summary_metadata["total_concepts"] >= 1
        assert result.summary_metadata["total_learning_objectives"] >= 1
    finally:
        if os.path.exists(temp_path):
            os.remove(temp_path)


SAMPLE_ENGLISH_SGK_BOOKMAP = """
Book map
ME AND MY FRIENDS

Unit 1 My friends Page 10
Competences: Asking and answering questions about where you are from
Structures: Where are you from? - I'm from...
Vocabulary: America, Australia, Britain, Japan, Malaysia, Singapore, Thailand, Viet Nam
Phonics: America, Australia

Unit 2 Time and daily routines Page 16
Competences: Asking and telling the time
Structures: What time is it? - It's...
Vocabulary: at, fifteen, forty-five, o'clock, thirty, get up, go to bed, go to school, have breakfast
Phonics: get, bed

Unit 3 My week Page 22
Competences: Asking and answering questions about the days of the week
Structures: What day is it today? - It's...
Vocabulary: Friday, Monday, Saturday, Sunday, Thursday, Tuesday, Wednesday
Phonics: music, Sunday
"""


def test_english_sgk_structure_extractor():
    with tempfile.NamedTemporaryFile("w", encoding="utf-8", suffix=".txt", delete=False) as f:
        f.write(SAMPLE_ENGLISH_SGK_BOOKMAP)
        temp_path = f.name

    try:
        parsed = TXTParser.parse(temp_path)
        extractor = StructureExtractor(
            subject="Tiếng Anh",
            grade=5,
            book_series="Global Success",
            title="Sách Giáo Khoa Tiếng Anh 5 Tập 1",
        )
        result = extractor.extract(parsed)

        # Verify Chunks & Unit extraction
        assert len(result.chunks) > 0
        units_found = [c.lesson for c in result.chunks if c.lesson and "Unit" in c.lesson]
        assert len(units_found) >= 1, "Cần trích xuất được các Unit tiếng Anh"

        # Check 8-level tree structure
        assert len(result.root_nodes) == 1
        subject_node = result.root_nodes[0]
        assert subject_node.title == "Tiếng Anh"

        book_node = subject_node.children[0].children[0]
        assert len(book_node.children) >= 1

        all_lessons = [les.title for chap in book_node.children for les in chap.children]
        assert any("Unit 1" in t for t in all_lessons), f"Không thấy Unit 1 trong danh sách bài học: {all_lessons}"
        assert any("Unit 2" in t for t in all_lessons), f"Không thấy Unit 2 trong danh sách bài học: {all_lessons}"
        assert any("Unit 3" in t for t in all_lessons), f"Không thấy Unit 3 trong danh sách bài học: {all_lessons}"
    finally:
        if os.path.exists(temp_path):
            os.remove(temp_path)
