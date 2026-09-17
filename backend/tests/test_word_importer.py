import io
import docx
import pytest
from fastapi import status
from app.services.word_importer import WordQuestionParser, import_word_questions


def create_sample_docx() -> bytes:
    """Create an in-memory .docx document containing sample test questions."""
    doc = docx.Document()

    # Question 1: Standard format with explicit answer and explanation
    doc.add_paragraph("Câu 1: Phép tính nào sau đây có kết quả bằng 100?")
    doc.add_paragraph("A. 25 x 3")
    doc.add_paragraph("B. 20 x 5")
    doc.add_paragraph("C. 15 x 6")
    doc.add_paragraph("D. 12 x 8")
    doc.add_paragraph("Đáp án: B")
    doc.add_paragraph("Lời giải: Ta có 20 x 5 = 100. Do đó chọn đáp án B.")

    # Question 2: Bold option format without separate "Đáp án:" line
    doc.add_paragraph("Câu 2. Trong câu sau, từ nào là danh từ riêng?")
    doc.add_paragraph("A. sông")
    p_b = doc.add_paragraph()
    run_b = p_b.add_run("B. Đà Lạt")
    run_b.bold = True
    doc.add_paragraph("C. thành phố")
    doc.add_paragraph("D. hoa")
    doc.add_paragraph("Giải thích: Đà Lạt là danh từ riêng chỉ địa danh.")

    # Question 3: Inline options
    doc.add_paragraph("Câu 3: 1 km bằng bao nhiêu mét?")
    doc.add_paragraph("A. 10 m\tB. 100 m\tC. 1000 m\tD. 10000 m")
    doc.add_paragraph("Đáp án: C")

    # Question 4: Malformed question (missing answer)
    doc.add_paragraph("Câu 4: Đây là câu hỏi lỗi không có đáp án đúng")
    doc.add_paragraph("A. Phương án 1")
    doc.add_paragraph("B. Phương án 2")
    doc.add_paragraph("C. Phương án 3")
    doc.add_paragraph("D. Phương án 4")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_word_parser_extracts_questions():
    docx_bytes = create_sample_docx()
    parser = WordQuestionParser(docx_bytes)
    result = parser.parse()

    assert result.total_detected == 4
    assert result.valid_count == 3
    assert result.invalid_count == 1
    assert len(result.errors) == 1
    assert "Câu 4" in result.errors[0]

    # Verify Question 1
    q1 = result.questions[0]
    assert q1.question_index == 1
    assert "Phép tính nào sau đây" in q1.content
    assert len(q1.options) == 4
    assert q1.correct_option == "B"
    assert q1.explanation is not None and "20 x 5 = 100" in q1.explanation

    # Verify Question 2 (Bold option)
    q2 = result.questions[1]
    assert q2.question_index == 2
    assert q2.correct_option == "B"
    assert q2.is_valid is True

    # Verify Question 3 (Inline)
    q3 = result.questions[2]
    assert q3.question_index == 3
    assert q3.correct_option == "C"
    assert q3.is_valid is True


def test_import_word_questions_to_db(db, teacher_user):
    docx_bytes = create_sample_docx()
    result = import_word_questions(
        db=db,
        docx_bytes=docx_bytes,
        user_id=teacher_user.id,
        subject="Toán",
        grade=4,
        chapter="Chương 1",
        lesson="Bài 1",
        commit=True,
    )

    assert result.total_detected == 4
    assert result.valid_count == 3
    assert result.imported_count == 3


def test_api_import_word_preview_and_commit(client, teacher_headers):
    docx_bytes = create_sample_docx()

    # 1. Preview mode (commit=False)
    preview_res = client.post(
        "/api/v1/questions/import-word",
        data={
            "subject": "Toán",
            "grade": 5,
            "commit": "false",
        },
        files={
            "file": ("test_exam.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
        },
        headers=teacher_headers,
    )
    assert preview_res.status_code == status.HTTP_200_OK
    preview_data = preview_res.json()
    assert preview_data["total_detected"] == 4
    assert preview_data["valid_count"] == 3
    assert preview_data["invalid_count"] == 1
    assert preview_data["imported_count"] == 0

    # 2. Commit mode (commit=True)
    commit_res = client.post(
        "/api/v1/questions/import-word",
        data={
            "subject": "Toán",
            "grade": 5,
            "commit": "true",
        },
        files={
            "file": ("test_exam.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
        },
        headers=teacher_headers,
    )
    assert commit_res.status_code == status.HTTP_200_OK
    commit_data = commit_res.json()
    assert commit_data["imported_count"] == 3

    # Check that questions are now in the question bank
    list_res = client.get("/api/v1/questions/?subject=Toán&grade=5", headers=teacher_headers)
    assert list_res.status_code == status.HTTP_200_OK
    assert list_res.json()["total"] >= 3
