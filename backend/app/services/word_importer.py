import io
import re
from typing import BinaryIO, List, Optional, Tuple, Union
import docx
from sqlalchemy.orm import Session

from app.models.question import (
    Question,
    QuestionDifficulty,
    QuestionOption,
    QuestionSource,
    QuestionStatus,
    QuestionType,
)
from app.schemas.question import (
    ParsedQuestionItem,
    QuestionOptionCreate,
    WordImportResult,
)


class WordQuestionParser:
    """Parser for Vietnamese question bank documents (.docx)."""

    QUESTION_HEADER_PATTERN = re.compile(
        r"^(?:Câu|Bài|Question)\s*(\d+)[\s.:\-–—]\s*(.*)$",
        re.IGNORECASE,
    )
    OPTION_LINE_PATTERN = re.compile(
        r"^\s*([A-D])[\s.)\-–—]\s*(.*)$"
    )
    INLINE_OPTION_PATTERN = re.compile(
        r"(?:^|\s{2,}|\t)([A-D])[\s.)\-–—]\s*(.+?)(?=(?:\s{2,}|\t)[A-D][\s.)\-–—]|$)"
    )
    ANSWER_PATTERN = re.compile(
        r"^\s*(?:\*?\s*(?:Đáp án|Đ/A|Answer|Key)[\s.:\-–—]\s*([A-D]))",
        re.IGNORECASE,
    )
    EXPLANATION_PATTERN = re.compile(
        r"^\s*(?:\*?\s*(?:Lời giải|Giải thích|Hướng dẫn giải|Explanation)[\s.:\-–—]\s*(.*))",
        re.IGNORECASE,
    )

    def __init__(self, docx_source: Union[BinaryIO, bytes, str]):
        if isinstance(docx_source, bytes):
            self.doc = docx.Document(io.BytesIO(docx_source))
        elif hasattr(docx_source, "read"):
            self.doc = docx.Document(docx_source)
        else:
            self.doc = docx.Document(str(docx_source))

    @staticmethod
    def _extract_options_from_line(text: str) -> Optional[dict[str, str]]:
        """Extract options A, B, C, D whether on single lines or multiple inline."""
        matches = list(re.finditer(r"(?:^|[\s\t]+)([A-D])[\s.)\-–—]\s*", text))
        if len(matches) >= 2:
            keys = [m.group(1).upper() for m in matches]
            # Ensure keys are unique
            if len(keys) == len(set(keys)):
                opts: dict[str, str] = {}
                for i in range(len(matches)):
                    curr_key = keys[i]
                    start_content = matches[i].end()
                    end_content = matches[i + 1].start() if i + 1 < len(matches) else len(text)
                    opts[curr_key] = text[start_content:end_content].strip()
                return opts
        elif len(matches) == 1 and matches[0].start() == 0:
            key = matches[0].group(1).upper()
            content = text[matches[0].end():].strip()
            return {key: content}
        return None

    def _extract_text_blocks(self) -> List[Tuple[str, bool]]:
        """Extract all paragraph texts and track if the paragraph has bold text runs."""
        blocks: List[Tuple[str, bool]] = []
        for p in self.doc.paragraphs:
            text = p.text.strip()
            if not text:
                continue
            is_bold = any(run.bold for run in p.runs if run.bold is not None)
            blocks.append((text, is_bold))

        # Also inspect tables if any
        for table in self.doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        text = p.text.strip()
                        if text:
                            is_bold = any(run.bold for run in p.runs if run.bold is not None)
                            blocks.append((text, is_bold))
        return blocks

    def parse(self) -> WordImportResult:
        """Parse all questions from the Word document."""
        blocks = self._extract_text_blocks()
        raw_questions: List[List[Tuple[str, bool]]] = []
        current_chunk: List[Tuple[str, bool]] = []

        for text, is_bold in blocks:
            header_match = self.QUESTION_HEADER_PATTERN.match(text)
            if header_match:
                if current_chunk:
                    raw_questions.append(current_chunk)
                current_chunk = [(text, is_bold)]
            else:
                if current_chunk:
                    current_chunk.append((text, is_bold))

        if current_chunk:
            raw_questions.append(current_chunk)

        parsed_items: List[ParsedQuestionItem] = []
        errors: List[str] = []

        for idx, chunk in enumerate(raw_questions, start=1):
            item, err = self._parse_single_question(idx, chunk)
            parsed_items.append(item)
            if err:
                errors.append(err)

        valid_count = sum(1 for q in parsed_items if q.is_valid)
        invalid_count = len(parsed_items) - valid_count

        return WordImportResult(
            total_detected=len(parsed_items),
            valid_count=valid_count,
            invalid_count=invalid_count,
            imported_count=0,
            questions=parsed_items,
            errors=errors,
        )

    def _parse_single_question(
        self, fallback_index: int, chunk: List[Tuple[str, bool]]
    ) -> Tuple[ParsedQuestionItem, Optional[str]]:
        header_text, _ = chunk[0]
        match = self.QUESTION_HEADER_PATTERN.match(header_text)
        q_idx = int(match.group(1)) if match else fallback_index
        first_line_content = match.group(2).strip() if match else header_text

        content_parts: List[str] = []
        if first_line_content:
            content_parts.append(first_line_content)

        options_dict: dict[str, str] = {}
        bold_options: set[str] = set()
        explicit_answer: Optional[str] = None
        explanation_parts: List[str] = []

        mode = "CONTENT"  # CONTENT -> OPTIONS -> EXPLANATION

        for text, is_bold in chunk[1:]:
            ans_match = self.ANSWER_PATTERN.match(text)
            if ans_match:
                explicit_answer = ans_match.group(1).upper()
                mode = "ANSWER"
                continue

            exp_match = self.EXPLANATION_PATTERN.match(text)
            if exp_match:
                mode = "EXPLANATION"
                exp_text = exp_match.group(1).strip()
                if exp_text:
                    explanation_parts.append(exp_text)
                continue

            if mode == "EXPLANATION":
                explanation_parts.append(text)
                continue

            # Extract options (handles both inline A/B/C/D and standalone A./B./C./D.)
            extracted_opts = self._extract_options_from_line(text)
            if extracted_opts:
                mode = "OPTIONS"
                options_dict.update(extracted_opts)
                if len(extracted_opts) == 1 and is_bold:
                    bold_options.add(list(extracted_opts.keys())[0])
                continue

            if mode == "CONTENT":
                content_parts.append(text)
            elif mode == "OPTIONS":
                # If continuing options text
                if options_dict:
                    last_key = list(options_dict.keys())[-1]
                    options_dict[last_key] += f" {text}"
                else:
                    content_parts.append(text)

        final_content = "\n".join(content_parts).strip()
        explanation = "\n".join(explanation_parts).strip() or None

        # Determine correct option
        correct_key: Optional[str] = explicit_answer
        if not correct_key and len(bold_options) == 1:
            correct_key = list(bold_options)[0]

        # Validation checks
        error_msg: Optional[str] = None
        if not final_content:
            error_msg = f"Câu {q_idx}: Nội dung câu hỏi trống."
        elif len(options_dict) < 2:
            error_msg = f"Câu {q_idx}: Tìm thấy ít hơn 2 phương án ({len(options_dict)} phương án)."
        elif not correct_key:
            error_msg = f"Câu {q_idx}: Không tìm thấy đáp án đúng (thiếu dòng 'Đáp án: [A-D]' hoặc phương án in đậm)."
        elif correct_key not in options_dict:
            error_msg = f"Câu {q_idx}: Đáp án đúng '{correct_key}' không nằm trong các phương án {list(options_dict.keys())}."

        is_valid = error_msg is None

        # Build options list
        options_list: List[QuestionOptionCreate] = []
        ordered_keys = ["A", "B", "C", "D"]
        # If standard keys exist, use them in order, else all keys found
        all_keys = [k for k in ordered_keys if k in options_dict] or sorted(options_dict.keys())
        for idx_opt, key in enumerate(all_keys):
            options_list.append(
                QuestionOptionCreate(
                    option_key=key,
                    content=options_dict.get(key, ""),
                    is_correct=(key == correct_key),
                    explanation=None,
                    order_index=idx_opt,
                )
            )

        item = ParsedQuestionItem(
            question_index=q_idx,
            raw_header=header_text,
            content=final_content,
            options=options_list,
            correct_option=correct_key,
            explanation=explanation,
            is_valid=is_valid,
            error_message=error_msg,
        )

        return item, error_msg


def import_word_questions(
    db: Session,
    docx_bytes: bytes,
    user_id: int,
    subject: str,
    grade: int,
    chapter: Optional[str] = None,
    lesson: Optional[str] = None,
    difficulty: QuestionDifficulty = QuestionDifficulty.MEDIUM,
    initial_status: QuestionStatus = QuestionStatus.REVIEW,
    commit: bool = True,
) -> WordImportResult:
    """Parse a .docx file and optionally persist valid questions to the database."""
    parser = WordQuestionParser(docx_bytes)
    result = parser.parse()

    if not commit or result.valid_count == 0:
        return result

    imported_count = 0
    for q_data in result.questions:
        if not q_data.is_valid:
            continue

        question = Question(
            content=q_data.content,
            question_type=QuestionType.MULTIPLE_CHOICE_SINGLE,
            difficulty=difficulty,
            status=initial_status,
            source=QuestionSource.WORD_IMPORT,
            subject=subject,
            grade=grade,
            chapter=chapter,
            lesson=lesson,
            explanation=q_data.explanation,
            created_by_id=user_id,
        )
        db.add(question)
        db.flush()  # To populate question.id

        for opt in q_data.options:
            db_opt = QuestionOption(
                question_id=question.id,
                option_key=opt.option_key,
                content=opt.content,
                is_correct=opt.is_correct,
                explanation=opt.explanation,
                order_index=opt.order_index,
            )
            db.add(db_opt)

        imported_count += 1

    db.commit()
    result.imported_count = imported_count
    return result
