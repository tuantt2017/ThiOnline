import os
import re
from dataclasses import dataclass, field
from typing import List, Optional
import pymupdf
import docx

from app.models.document import DocumentType


class DocumentParsingError(Exception):
    """Raised when parsing a document fails."""
    pass


@dataclass
class ParsedPage:
    page_number: int
    text: str
    headings: List[str] = field(default_factory=list)


@dataclass
class ParsedDocument:
    filename: str
    file_type: DocumentType
    pages: List[ParsedPage]
    total_pages: int
    total_chars: int
    raw_text: str


class PDFParser:
    @staticmethod
    def parse(file_path: str) -> ParsedDocument:
        if not os.path.exists(file_path):
            raise DocumentParsingError(f"Tệp không tồn tại: {file_path}")

        try:
            doc = pymupdf.open(file_path)
        except Exception as e:
            raise DocumentParsingError(f"Không thể mở tệp PDF: {str(e)}")

        if doc.is_encrypted:
            raise DocumentParsingError("Tệp PDF đã bị khóa bằng mật khẩu. Vui lòng mở khóa trước khi tải lên.")

        if len(doc) == 0:
            raise DocumentParsingError("Tệp PDF rỗng, không chứa trang nào.")

        pages: List[ParsedPage] = []
        full_text_list: List[str] = []

        for page_idx in range(len(doc)):
            page = doc[page_idx]
            page_text = page.get_text("text").strip()
            
            # Detect basic headings from text blocks if available
            headings: List[str] = []
            blocks = page.get_text("blocks")
            for block in blocks:
                # block: (x0, y0, x1, y1, text, block_no, block_type)
                if len(block) >= 5 and block[6] == 0:  # text block
                    line = block[4].strip()
                    if re.match(r"^(Chương|CHƯƠNG|Bài|BÀI|Chủ đề|CHỦ ĐỀ|Phần|PHẦN)\b", line, re.IGNORECASE):
                        headings.append(line)

            pages.append(
                ParsedPage(
                    page_number=page_idx + 1,
                    text=page_text,
                    headings=headings,
                )
            )
            if page_text:
                full_text_list.append(page_text)

        raw_text = "\n\n".join(full_text_list)
        if not raw_text.strip():
            raise DocumentParsingError("Tệp PDF không chứa văn bản có thể trích xuất (có thể là tệp scan ảnh thuần túy).")

        return ParsedDocument(
            filename=os.path.basename(file_path),
            file_type=DocumentType.PDF,
            pages=pages,
            total_pages=len(pages),
            total_chars=len(raw_text),
            raw_text=raw_text,
        )


class DOCXParser:
    @staticmethod
    def parse(file_path: str) -> ParsedDocument:
        if not os.path.exists(file_path):
            raise DocumentParsingError(f"Tệp không tồn tại: {file_path}")

        try:
            doc = docx.Document(file_path)
        except Exception as e:
            raise DocumentParsingError(f"Không thể đọc tệp DOCX: {str(e)}")

        paragraphs_text: List[str] = []
        headings: List[str] = []

        for p in doc.paragraphs:
            text = p.text.strip()
            if not text:
                continue
            paragraphs_text.append(text)
            # Check style name for headings
            style_name = getattr(p.style, "name", "")
            if "Heading" in style_name or re.match(r"^(Chương|CHƯƠNG|Bài|BÀI|Chủ đề|Mục)\b", text, re.IGNORECASE):
                headings.append(text)

        # Also extract text from tables
        for table in doc.tables:
            for row in table.rows:
                row_texts = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_texts:
                    paragraphs_text.append(" | ".join(row_texts))

        raw_text = "\n\n".join(paragraphs_text)
        if not raw_text.strip():
            raise DocumentParsingError("Tệp DOCX rỗng hoặc không chứa văn bản.")

        # In DOCX there are no native page numbers unless rendered;
        # We chunk into virtual pages of approximately 500 words (~3000 chars) for navigation
        PAGE_CHAR_SIZE = 3000
        pages: List[ParsedPage] = []
        chunks = [raw_text[i:i + PAGE_CHAR_SIZE] for i in range(0, len(raw_text), PAGE_CHAR_SIZE)]
        for idx, chunk_text in enumerate(chunks):
            page_headings = [h for h in headings if h in chunk_text]
            pages.append(
                ParsedPage(
                    page_number=idx + 1,
                    text=chunk_text.strip(),
                    headings=page_headings,
                )
            )

        return ParsedDocument(
            filename=os.path.basename(file_path),
            file_type=DocumentType.DOCX,
            pages=pages,
            total_pages=len(pages),
            total_chars=len(raw_text),
            raw_text=raw_text,
        )


class TXTParser:
    @staticmethod
    def parse(file_path: str) -> ParsedDocument:
        if not os.path.exists(file_path):
            raise DocumentParsingError(f"Tệp không tồn tại: {file_path}")

        encodings = ["utf-8", "utf-8-sig", "utf-16", "cp1258", "latin-1"]
        raw_text = None

        for enc in encodings:
            try:
                with open(file_path, "r", encoding=enc) as f:
                    raw_text = f.read()
                break
            except (UnicodeDecodeError, UnicodeError):
                continue

        if raw_text is None:
            raise DocumentParsingError("Không thể giải mã tệp văn bản. Vui lòng lưu tệp ở định dạng UTF-8.")

        raw_text = raw_text.strip()
        if not raw_text:
            raise DocumentParsingError("Tệp văn bản rỗng.")

        PAGE_CHAR_SIZE = 3000
        pages: List[ParsedPage] = []
        chunks = [raw_text[i:i + PAGE_CHAR_SIZE] for i in range(0, len(raw_text), PAGE_CHAR_SIZE)]
        for idx, chunk_text in enumerate(chunks):
            pages.append(
                ParsedPage(
                    page_number=idx + 1,
                    text=chunk_text.strip(),
                    headings=[],
                )
            )

        return ParsedDocument(
            filename=os.path.basename(file_path),
            file_type=DocumentType.TXT,
            pages=pages,
            total_pages=len(pages),
            total_chars=len(raw_text),
            raw_text=raw_text,
        )


def parse_document_file(file_path: str, file_type: DocumentType) -> ParsedDocument:
    """Parse document according to its file type."""
    if file_type == DocumentType.PDF:
        return PDFParser.parse(file_path)
    elif file_type == DocumentType.DOCX:
        return DOCXParser.parse(file_path)
    elif file_type == DocumentType.TXT:
        return TXTParser.parse(file_path)
    else:
        raise DocumentParsingError(f"Loại tài liệu không được hỗ trợ: {file_type}")
